from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# Initialize presentation with 16:9 widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_layout)

# Color Palette: Slate & Deep Corporate Blue
COLOR_BG = RGBColor(248, 250, 252)         # Slate 50
COLOR_TITLE = RGBColor(15, 23, 42)         # Slate 900
COLOR_SUBTITLE = RGBColor(37, 99, 235)     # Blue 600
COLOR_BOX_BG = RGBColor(255, 255, 255)     # Clean white card
COLOR_BORDER = RGBColor(203, 213, 225)     # Slate 300
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)     # Slate 800
COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
COLOR_ARROW = RGBColor(59, 130, 246)       # Blue 500

# Set solid background
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = COLOR_BG

# --- Header Section ---
header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
tf = header_box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p_sub = tf.paragraphs[0]
p_sub.text = "SYSTEM ARCHITECTURE & PIPELINE"
p_sub.font.size = Pt(11)
p_sub.font.bold = True
p_sub.font.color.rgb = COLOR_SUBTITLE
p_sub.space_after = Pt(4)

p_main = tf.add_paragraph()
p_main.text = "End-to-End Processing Workflow"
p_main.font.size = Pt(24)
p_main.font.bold = True
p_main.font.color.rgb = COLOR_TITLE

# --- Architecture Stages Data ---
stages = [
    {
        "badge": "STAGE 01",
        "title": "Streamlit UI",
        "desc": "Captures raw student text, task selection, and manages state."
    },
    {
        "badge": "STAGE 02",
        "title": "Guardrails",
        "desc": "Validates boundaries:\n10 <= length <= 8,000.\nStops zero-token calls."
    },
    {
        "badge": "STAGE 03",
        "title": "Prompt Engine",
        "desc": "Injects R-T-C-F rules, role constraints, and strict schema format."
    },
    {
        "badge": "STAGE 04",
        "title": "Groq LPU",
        "desc": "Llama 3.3 70B inference with fallback to 8B on timeouts/503s."
    },
    {
        "badge": "STAGE 05",
        "title": "SSE Streaming",
        "desc": "Streams chunks directly via st.write_stream (TTFT < 500ms)."
    }
]

# Layout dimensions (5 nodes + 4 arrows horizontally)
start_x = 0.8
box_width = 1.85
box_height = 3.6
arrow_width = 0.45
arrow_height = 0.35
box_y = 2.4
arrow_y = box_y + (box_height / 2.0) - (arrow_height / 2.0)

for idx, stage in enumerate(stages):
    current_x = start_x + idx * (box_width + arrow_width + 0.15)
    
    # 1. Draw Process Box (Card)
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(current_x),
        Inches(box_y),
        Inches(box_width),
        Inches(box_height)
    )
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_BOX_BG
    box.line.color.rgb = COLOR_BORDER
    box.line.width = Pt(1.5)
    
    btf = box.text_frame
    btf.word_wrap = True
    btf.margin_left = Inches(0.15)
    btf.margin_right = Inches(0.15)
    btf.margin_top = Inches(0.25)
    btf.margin_bottom = Inches(0.15)
    
    # Stage Badge
    p0 = btf.paragraphs[0]
    p0.text = stage["badge"]
    p0.alignment = PP_ALIGN.CENTER
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_SUBTITLE
    p0.space_after = Pt(8)
    
    # Stage Title
    p1 = btf.add_paragraph()
    p1.text = stage["title"]
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_MAIN
    p1.space_after = Pt(10)
    
    # Stage Details
    p2 = btf.add_paragraph()
    p2.text = stage["desc"]
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    
    # 2. Draw Connector Arrow (between nodes)
    if idx < len(stages) - 1:
        arrow_x = current_x + box_width + 0.075
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW,
            Inches(arrow_x),
            Inches(arrow_y),
            Inches(arrow_width),
            Inches(arrow_height)
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = COLOR_ARROW
        arrow.line.fill.background()  # Borderless arrow

# --- Bottom Callout / Annotation Card ---
note_box = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE,
    Inches(0.8),
    Inches(6.3),
    Inches(11.7),
    Inches(0.7)
)
note_box.fill.solid()
note_box.fill.fore_color.rgb = RGBColor(239, 246, 255) # Light blue accent
note_box.line.color.rgb = RGBColor(191, 219, 254)
note_box.line.width = Pt(1)

ntf = note_box.text_frame
ntf.word_wrap = True
ntf.margin_top = Inches(0.12)
np = ntf.paragraphs[0]
np.text = "Key Engineering Takeaway: Fault isolation between Stage 02 (Validation) and Stage 04 (Inference) ensures zero API cost on bad inputs, while Stage 05 streams partial tokens for real-time responsiveness."
np.font.size = Pt(11)
np.font.color.rgb = RGBColor(30, 58, 138)
np.font.bold = False

# Save deck
prs.save("Pipeline_Slide_Diagram.pptx")
print("Saved 'Pipeline_Slide_Diagram.pptx' successfully.")